#!/usr/bin/env python3
"""Turn review CSV into an Excel workbook with dropdowns/filters for the first reviewer."""
from __future__ import annotations
import csv, json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "docs" / "kg-review"
OUT = SRC / "JTE-图谱映射审核表.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1B5E7A")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN = Side(style="thin", color="D9E2E8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def load(name: str) -> list[dict]:
    with (SRC / name).open(encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def style_sheet(ws, widths: dict[str, int]) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def main() -> int:
    items = load("JTE-条目候选表.csv")
    rows = load("JTE-图谱映射审核表.csv")
    summary = json.loads((SRC / "summary.json").read_text(encoding="utf-8"))
    wb = Workbook()

    info = wb.active
    info.title = "说明"
    lines = [
        ("知识图谱 v2 × JTE 66 项映射审核表", True),
        ("", False),
        ("用途：把「评估项/不符合情形」映射到 v2 图谱候选故障路径，供报告检索使用。", False),
        ("当前状态：脚本自动生成候选（candidate），尚未进入任何报告。", False),
        ("", False),
        ("审核流程（两级）", True),
        ("1. 本人初审：在「情形审核」表填写 self_decision，批准有明确依据的候选；", False),
        ("   批准后可在报告附页作为「本人初审，待专家复核」的候选证据。", False),
        ("2. 专家复核：填写 expert_decision，升级为 expert_reviewed 后才作为正式候选证据。", False),
        ("", False),
        ("填写方法", True),
        ("self_decision 可选：批准 / 驳回 / 无合适候选 / 待定（留空=未审）。", False),
        ("同一情形有多条候选时，只给真正匹配的那一行填「批准」；其余可留空或「驳回」。", False),
        ("若候选都不对但图谱里另有路径，可在 self_note 填正确 path_id（从「条目候选」表或图谱查）。", False),
        ("「无合适候选」= 该情形在图谱中确实没有可用路径，运行时转自由文本检索/规则处置。", False),
        ("", False),
        ("自动候选覆盖", True),
        (f"条目：{summary['items_with_candidates']}/{summary['items']} 有候选；"
         f"情形：{summary['options_with_candidates']}/{summary['options']} 有候选；"
         f"覆盖路径：{summary['candidate_paths_covered']} 条。", False),
        ("候选由名称/别名/字符相似度生成，质量参差；驳回是正常操作，不要为了覆盖率勉强批准。", False),
        ("", False),
        ("文件与版本", True),
        ("图谱：v2_1000paths（1,052 条候选路径，confidence=0.4，verified=false）", False),
        ("本表生成：tools/kg/build_mapping_candidates.py；索引：backend/data/kg_v2.sqlite", False),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        cell = info.cell(row=i, column=1, value=text)
        cell.font = Font(bold=bold, size=13 if i == 1 else 10.5,
                         color="1B5E7A" if bold else "333333")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    info.column_dimensions["A"].width = 110
    info.sheet_view.showGridLines = False

    ws_items = wb.create_sheet("条目候选")
    item_cols = ["jte_no", "jte_name", "graph_system", "graph_component", "component_id",
                 "component_score", "candidate_path_count", "sample_path_ids",
                 "self_decision", "self_note", "expert_decision", "expert_note"]
    ws_items.append(item_cols)
    for row in items:
        ws_items.append([row.get(c, "") for c in item_cols])
    style_sheet(ws_items, {"A": 9, "B": 26, "C": 16, "D": 30, "E": 34, "F": 10,
                           "G": 9, "H": 70, "I": 12, "J": 30, "K": 12, "L": 30})

    ws = wb.create_sheet("情形审核")
    cols = ["jte_no", "jte_name", "option_key", "option_text", "option_category",
            "rule_measure", "rank", "score", "path_id", "system", "component",
            "path_summary", "inspection", "action", "verification", "risk",
            "match_reason", "self_decision", "self_note", "expert_decision", "expert_note"]
    ws.append(cols)
    for row in rows:
        ws.append([row.get(c, "") for c in cols])
    style_sheet(ws, {"A": 8, "B": 24, "C": 9, "D": 55, "E": 8, "F": 45, "G": 6,
                     "H": 8, "I": 40, "J": 14, "K": 30, "L": 55, "M": 40, "N": 40,
                     "O": 40, "P": 30, "Q": 28, "R": 12, "S": 30, "T": 12, "U": 30})
    dv = DataValidation(type="list", formula1='"批准,驳回,无合适候选,待定"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"R2:R{ws.max_row}")
    dv2 = DataValidation(type="list", formula1='"批准,驳回,待定"', allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add(f"I2:I{ws_items.max_row}")
    dv2.add(f"T2:T{ws.max_row}")
    green = PatternFill("solid", fgColor="E8F5E9")
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(row=r, column=7).value) == "1" and ws.cell(row=r, column=9).value:
            ws.cell(row=r, column=7).fill = green

    wb.save(OUT)
    print(str(OUT))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
